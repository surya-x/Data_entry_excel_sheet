from pdfminer.pdfinterp import PDFResourceManager, PDFPageInterpreter
from pdfminer.converter import TextConverter
from pdfminer.layout import LAParams
from pdfminer.pdfpage import PDFPage
from pdfminer.pdfinterp import resolve1
from pdfminer.pdfparser import PDFParser
from pdfminer.pdfdocument import PDFDocument

import time
import io
import openpyxl
import sys
from config import *


def get_text_from_pdf(path):
    logging.info("method 'get_text_from_pdf' called ")
    data = []
    try:

        fp = open(path, 'rb')
        parser = PDFParser(fp)
        doc = PDFDocument(parser)
        max_pages = resolve1(doc.catalog['Pages'])['Count']

        for i in range(0, max_pages, 2):
            data.append(convert_pdf_to_txt(fp, i))

    except Exception as e:
        print("Error : Reading pages from pdf")
        print(e)
        logging.error(e)
        sys.exit()
    finally:
        fp.close()
        return data


def convert_pdf_to_txt(fp, pageNumber):
    logging.info("method 'convert_pdf_to_txt' called ")
    try:
        rsrcmgr = PDFResourceManager()
        retstr = io.StringIO()
        laparams = LAParams()
        device = TextConverter(
            rsrcmgr, retstr, codec='utf-8', laparams=laparams)

        interpreter = PDFPageInterpreter(rsrcmgr, device)
        password = ""
        maxpages = 0
        caching = False
        pagenos = set()

        for pageN, page in enumerate(PDFPage.get_pages(fp, pagenos, maxpages=maxpages,
                                                       password=password,
                                                       caching=caching,
                                                       check_extractable=True)):
            if pageN == pageNumber:
                interpreter.process_page(page)

        device.close()
        text = retstr.getvalue()
        retstr.close()
        return text
    except Exception as e:
        print("Error : Converting pages from pdf")
        print(e)
        logging.error(e)
        sys.exit()


def write_data(excel_path, excel_format, data):
    logging.info("method write_data called ")
    try:
        workbook = openpyxl.load_workbook(excel_format)
        sheet = workbook.worksheets[0]

        for index, each_row in enumerate(data):
            for i, each_cell in zip(range(1, 13), each_row):
                sheet.cell(row=index + 2, column=i).value = each_cell

        workbook.save(excel_path)
    except PermissionError:
        print("Error Occured : Close the Excel Window before executing code")
        sys.exit()
    except Exception as e:
        print("Error : Writing in full_pay.xlsx")
        print(e)
        logging.error(e)
        sys.exit()
